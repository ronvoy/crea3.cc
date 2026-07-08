"""Excel (.xlsx) export of an allocation proposal.

Companion to report_pdf.build_report_pdf: the PDF is the polished narrative
document; this workbook is the same content in analysable, spreadsheet form so
the parties (or a mediator) can sort/filter the numbers. It carries:

  * Overview        - dispute reference, status, generation metadata.
  * Parties         - entitlement shares (assigned / claimed / used) + status.
  * Assets          - each asset, each party's valuation, the assigned party and
                      value, the valuation gap, divisible split fractions.
  * Settlement      - per-party value received, the equalising cash settlement.
  * Decisions       - each party's accept / reject decision on the proposal
                      (this is what makes the report reflect the outcome).

The numbers are exactly those produced by proposals_service.build_proposal and
stored on the AllocationProposal (outputs/metrics); nothing is recomputed here.
"""

from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- palette (kept close to the PDF's brand styling) -----------------------
BRAND = "1E3A5F"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
SUBHEAD_FILL = PatternFill("solid", fgColor="E8EEF5")
ACCEPT_FILL = PatternFill("solid", fgColor="E7F5EC")
REJECT_FILL = PatternFill("solid", fgColor="FBEAEA")
PENDING_FILL = PatternFill("solid", fgColor="FFF6E5")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
MUTED = Font(color="6B7280", size=9)
THIN = Side(style="thin", color="D5DCE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR_FMT = "#,##0.00\\ \u20ac"
PCT_FMT = "0.0%"


def _num(v: Any) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _title_row(ws, row: int, text: str, span: int) -> int:
    c = ws.cell(row=row, column=1, value=text)
    c.font = WHITE_BOLD
    c.alignment = Alignment(vertical="center")
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).fill = HEADER_FILL
    ws.row_dimensions[row].height = 20
    return row + 1


def _header_cells(ws, row: int, headers: list[str]) -> int:
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = BOLD
        c.fill = SUBHEAD_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return row + 1


def _row(ws, row: int, values: list[Any], *, money_cols: set[int] | None = None,
         pct_cols: set[int] | None = None, fills: dict[int, PatternFill] | None = None) -> int:
    money_cols = money_cols or set()
    pct_cols = pct_cols or set()
    fills = fills or {}
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.border = BORDER
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if col in money_cols and isinstance(v, (int, float)):
            c.number_format = EUR_FMT
        if col in pct_cols and isinstance(v, (int, float)):
            c.number_format = PCT_FMT
        if col in fills:
            c.fill = fills[col]
    return row + 1


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def build_report_xlsx(
    *,
    out_path: str,
    dispute: Any,
    agents: list[Any],
    goods: list[Any],
    proposal: Any,
    acceptances: list[dict[str, Any]] | None = None,
    kind: str = "proposal",
    lang: str = "en",  # reserved; labels are English (matches the PDF's generated text)
) -> str:
    outputs = dict(getattr(proposal, "outputs", {}) or {})
    metrics = dict(getattr(proposal, "metrics", {}) or {})
    allocations = outputs.get("allocations", []) or []
    acceptances = acceptances or []

    non_mediators = [a for a in agents if (getattr(a, "role_in_dispute", None) or "agent").lower() != "mediator"]
    name_by_id = {int(a.id): (a.name or f"Party #{a.id}") for a in non_mediators if a.id is not None}
    agent_ids = list(name_by_id.keys())

    comp = metrics.get("compensation_by_agent", {}) or {}
    subj = metrics.get("subjective_value_by_agent", {}) or {}
    perceived = metrics.get("perceived_total_by_agent", {}) or {}
    equalized = metrics.get("equalized_value_by_agent", {}) or {}
    ent = metrics.get("entitlement_by_agent", {}) or {}
    ent_detail = {int(d["agent_id"]): d for d in (metrics.get("entitlement_detail") or [])}

    wb = Workbook()

    # ---------------------------------------------------------------- Overview
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    r = 1
    r = _title_row(ws, r, "CREA3 — %s" % ("Dispute Resolution Proposal" if kind == "proposal" else "Final Dispute Resolution Report"), 2)
    gen = datetime.now(timezone.utc).strftime("%d %b %Y, %H:%M UTC")
    info = [
        ("Dispute", f"#{getattr(dispute, 'id', '')} — {getattr(dispute, 'title', '')}"),
        ("Document type", "Proposal" if kind == "proposal" else "Final report"),
        ("Status", str(getattr(dispute, "status", "") or "").capitalize()),
        ("Generated (UTC)", gen),
        ("Algorithm", str(metrics.get("algorithm", "") or getattr(proposal, "algorithm_version", ""))),
        ("Total estimated value", _num(metrics.get("total_estimated_value"))),
        ("Balancing note", str(metrics.get("compensation_note", "") or outputs.get("compensation_note", ""))),
    ]
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = BOLD
        c = ws.cell(row=r, column=2, value=v)
        if isinstance(v, (int, float)):
            c.number_format = EUR_FMT
        ws.cell(row=r, column=1).border = BORDER
        c.border = BORDER
        r += 1
    ws.cell(row=r + 1, column=1, value="Monetary values are those entered by the parties. Star ratings are preferences and never change an asset's value.").font = MUTED
    _autosize(ws, {1: 26, 2: 62})

    # ------------------------------------------------------------------ Parties
    ws = wb.create_sheet("Parties")
    ws.sheet_view.showGridLines = False
    r = 1
    r = _title_row(ws, r, "Parties to the dispute", 5)
    r = _header_cells(ws, r, ["Party", "Assigned share", "Claimed share", "Share used", "Invite status"])
    for a in non_mediators:
        d = ent_detail.get(int(a.id), {})
        claimed = d.get("claimed")
        r = _row(ws, r, [
            a.name or "—",
            _num(d.get("assigned", getattr(a, "entitlement_share", 0.0))),
            (_num(claimed) if claimed is not None else "—"),
            _num(d.get("normalized_share", ent.get(str(a.id), 0.0))),
            str(getattr(a, "invite_status", "") or "—").capitalize(),
        ], pct_cols={2, 3, 4})
    if bool(metrics.get("entitlement_mismatch")):
        ws.cell(row=r + 1, column=1,
                value="Note: the parties did not fully agree on entitlement shares; the shares used were normalised to sum to 100%.").font = MUTED
    _autosize(ws, {1: 28, 2: 16, 3: 16, 4: 14, 5: 16})

    # ------------------------------------------------------------------- Assets
    ws = wb.create_sheet("Assets & valuations")
    ws.sheet_view.showGridLines = False
    r = 1
    r = _title_row(ws, r, "Disputed assets, valuations and allocation", 6 + len(agent_ids))
    headers = ["Asset", "Estimated value"] + [f"Value — {name_by_id[a]}" for a in agent_ids] + \
              ["Assigned to", "Assigned value", "Valuation gap", "Notes"]
    r = _header_cells(ws, r, headers)
    money_cols = {2} | {3 + i for i in range(len(agent_ids))} | {3 + len(agent_ids) + 1, 3 + len(agent_ids) + 2}
    for al in allocations:
        pv = al.get("party_valuations", {}) or {}
        notes = []
        if al.get("divisible"):
            fr = al.get("fraction_by_name") or {}
            notes.append("Split: " + ", ".join(f"{k} {round(_num(v) * 100)}%" for k, v in fr.items()))
        if al.get("reconciled_to_mean"):
            notes.append("reconciled" + (f" (€{al.get('reconciled_value')})" if al.get("reconciled_value") is not None else ""))
        if al.get("contested_by_omission"):
            notes.append("omitted by: " + ", ".join(al.get("omitted_by_names") or []))
        row_vals = [al.get("good_name") or "—", _num(al.get("estimated_value"))]
        row_vals += [_num(pv.get(str(a))) for a in agent_ids]
        row_vals += [
            al.get("assigned_agent_name") or "—",
            _num(al.get("assigned_value")),
            _num(al.get("valuation_gap")),
            "; ".join(notes),
        ]
        r = _row(ws, r, row_vals, money_cols=money_cols)
    _autosize(ws, {1: 26, 2: 15, **{3 + i: 16 for i in range(len(agent_ids))},
                   3 + len(agent_ids): 20, 4 + len(agent_ids): 15, 5 + len(agent_ids): 14, 6 + len(agent_ids): 40})

    # --------------------------------------------------------------- Settlement
    ws = wb.create_sheet("Settlement")
    ws.sheet_view.showGridLines = False
    r = 1
    r = _title_row(ws, r, "Equitable settlement (Knaster sealed bids)", 5)
    r = _header_cells(ws, r, ["Party", "Perceived total", "Value received", "Cash settlement", "Equalised value"])
    for a in agent_ids:
        key = str(a)
        cash = _num(comp.get(key))
        r = _row(ws, r, [
            name_by_id[a],
            _num(perceived.get(key)),
            _num(subj.get(key)),
            cash,
            _num(equalized.get(key)),
        ], money_cols={2, 3, 4, 5})
    note = str(metrics.get("compensation_note", "") or "")
    if note:
        ws.cell(row=r + 1, column=1, value=note).font = MUTED
    ws.cell(row=r + 2, column=1, value="Cash settlement: negative = the party pays in, positive = the party receives.").font = MUTED
    _autosize(ws, {1: 26, 2: 16, 3: 16, 4: 16, 5: 16})

    # ---------------------------------------------------------------- Decisions
    ws = wb.create_sheet("Decisions")
    ws.sheet_view.showGridLines = False
    r = 1
    r = _title_row(ws, r, "Party decisions on the proposal", 3)
    r = _header_cells(ws, r, ["Party", "Decision", "Comment"])
    if acceptances:
        for a in acceptances:
            acc = a.get("accepted")
            label = "Accepted" if acc is True else ("Rejected" if acc is False else "Pending")
            fill = ACCEPT_FILL if acc is True else (REJECT_FILL if acc is False else PENDING_FILL)
            r = _row(ws, r, [a.get("name") or "—", label, a.get("comment") or ""],
                     fills={2: fill})
    else:
        r = _row(ws, r, ["—", "No decisions recorded yet", ""])
    _autosize(ws, {1: 26, 2: 16, 3: 52})

    wb.save(out_path)
    with open(out_path, "rb") as fh:
        return sha256(fh.read()).hexdigest()
