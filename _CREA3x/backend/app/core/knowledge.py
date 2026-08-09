"""Knowledge Base (RAG) core: parsing, chunking, indexing and retrieval.

Design goals for this fragile, low-dependency environment:
  * Heavy libraries (faiss, numpy, pypdf, python-docx, odfpy, rank_bm25) are
    imported LAZILY inside functions. A missing library degrades gracefully
    rather than breaking import of the whole app.
  * Retrieval prefers embeddings (FAISS/semantic) and automatically falls back
    to BM25 keyword search when embeddings are unavailable (no OpenRouter
    embeddings, or an error), so the KB always answers.

Public surface used by the admin API and the assistant:
  - PRESETS / default_params(preset)
  - extract_text(filename, content_type, data) -> str
  - add_document(session, section, filename, content_type, data) -> KbDocument
  - reindex(session, section) -> int              (chunks re-embedded)
  - retrieve(session, section, query, top_k=?) -> list[dict]
  - get_config / set_config / status
  - seed_workflow_doc(session)
"""

from __future__ import annotations

import json
import logging
import re
from datetime import datetime, timezone

from sqlmodel import Session, select

from . import embeddings
from ..models import KbDocument, KbChunk, KbIndexConfig, KB_SECTIONS

logger = logging.getLogger(__name__)


# ── Presets → retrieval/index parameters ─────────────────────────────────────
# chunk_size/overlap are used at upload time; top_k/min_score/temperature at
# query time. The admin can override any of these ("custom" preset).
PRESETS: dict[str, dict] = {
    "optimal":  {"top_k": 4, "min_score": 0.35, "temperature": 0.1, "chunk_size": 800,  "chunk_overlap": 100},
    "balanced": {"top_k": 6, "min_score": 0.20, "temperature": 0.35, "chunk_size": 1000, "chunk_overlap": 150},
    "creative": {"top_k": 8, "min_score": 0.10, "temperature": 0.7,  "chunk_size": 1200, "chunk_overlap": 200},
}
ENGINES = ("faiss", "bm25", "hybrid")


def default_params(preset: str) -> dict:
    return dict(PRESETS.get(preset, PRESETS["balanced"]))


def _now() -> datetime:
    return datetime.now(timezone.utc)


# ── Document text extraction ─────────────────────────────────────────────────
def extract_text(filename: str, content_type: str, data: bytes) -> str:
    """Best-effort plain-text extraction. Never raises: returns "" on failure."""
    name = (filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    try:
        if ext == "pdf" or "pdf" in (content_type or ""):
            return _extract_pdf(data)
        if ext == "docx" or "word" in (content_type or ""):
            return _extract_docx(data)
        if ext == "odt" or "opendocument" in (content_type or ""):
            return _extract_odt(data)
        if ext == "rtf" or "rtf" in (content_type or ""):
            return _extract_rtf(data)
        if ext == "json":
            try:
                return json.dumps(json.loads(data.decode("utf-8", "replace")), indent=2, ensure_ascii=False)
            except Exception:
                return data.decode("utf-8", "replace")
        # txt / md / csv / anything else: decode as text
        return data.decode("utf-8", "replace")
    except Exception as e:  # pragma: no cover - defensive
        logger.warning("extract_text failed for %s: %s", filename, e)
        return ""


def missing_parser(filename: str, content_type: str = "") -> str | None:
    """Return the pip package needed to read this file type if it's NOT installed.

    Lets callers tell "the server can't parse this type" apart from "the file has
    no extractable text". Returns None when the parser is available or the type
    needs no third-party library (txt/md/json/csv/rtf).
    """
    name = (filename or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    needed = {"pdf": ("pypdf", "pypdf"), "docx": ("docx", "python-docx"),
              "odt": ("odf", "odfpy")}
    entry = needed.get(ext)
    if not entry:
        if "pdf" in (content_type or ""):
            entry = ("pypdf", "pypdf")
        else:
            return None
    module, package = entry
    try:
        __import__(module)
        return None
    except Exception:
        return package


def _extract_pdf(data: bytes) -> str:
    import io
    from pypdf import PdfReader  # lazy; ImportError surfaces via missing_parser()
    reader = PdfReader(io.BytesIO(data))
    # Many real-world (esp. legal) PDFs are permission-encrypted with an empty
    # user password; pypdf then returns no text until decrypted.
    if getattr(reader, "is_encrypted", False):
        for pw in ("", "\x00"):
            try:
                if reader.decrypt(pw):
                    break
            except Exception:
                pass
    parts: list[str] = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception as e:
            logger.warning("PDF page extract failed: %s", e)
            txt = ""
        if txt.strip():
            parts.append(txt)
    return "\n\n".join(parts).strip()


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx, lazy
        import io
        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs).strip()
    except Exception as e:
        logger.warning("DOCX parse failed (is python-docx installed?): %s", e)
        return ""


def _extract_rtf(data: bytes) -> str:
    raw = data.decode("utf-8", "replace")
    try:
        from striprtf.striprtf import rtf_to_text  # lazy
        return (rtf_to_text(raw) or "").strip()
    except Exception:
        # Minimal fallback: drop RTF control words/groups.
        import re as _re
        txt = _re.sub(r"\\'[0-9a-fA-F]{2}", "", raw)
        txt = _re.sub(r"\\[a-zA-Z]+-?\d* ?", "", txt)
        txt = txt.replace("{", "").replace("}", "")
        return txt.strip()


def _extract_odt(data: bytes) -> str:
    try:
        from odf.opendocument import load  # odfpy, lazy
        from odf import text as odftext
        from odf.element import Element
        import io
        doc = load(io.BytesIO(data))
        parts: list[str] = []
        for para in doc.getElementsByType(odftext.P):
            parts.append(str(para))
        return "\n".join(parts).strip()
    except Exception as e:
        logger.warning("ODT parse failed (is odfpy installed?): %s", e)
        return ""


# ── Chunking ─────────────────────────────────────────────────────────────────
def chunk_text(text: str, size: int = 1000, overlap: int = 150) -> list[str]:
    """Split text into ~`size`-char chunks with `overlap`, on whitespace bounds."""
    text = re.sub(r"[ \t]+", " ", (text or "").strip())
    if not text:
        return []
    size = max(200, int(size))
    overlap = max(0, min(int(overlap), size - 50))
    chunks: list[str] = []
    start = 0
    n = len(text)
    while start < n:
        end = min(start + size, n)
        # Prefer to break on a whitespace boundary near `end`.
        if end < n:
            ws = text.rfind(" ", start + size - overlap, end)
            if ws > start:
                end = ws
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= n:
            break
        start = max(end - overlap, start + 1)
    return chunks


# ── Config ───────────────────────────────────────────────────────────────────
def get_config(session: Session, section: str) -> KbIndexConfig:
    cfg = session.exec(select(KbIndexConfig).where(KbIndexConfig.section == section)).first()
    if not cfg:
        cfg = KbIndexConfig(
            section=section, engine="faiss", preset="balanced",
            params_json=json.dumps(default_params("balanced")),
        )
        session.add(cfg)
        session.commit()
        session.refresh(cfg)
    return cfg


def config_params(cfg: KbIndexConfig) -> dict:
    try:
        p = json.loads(cfg.params_json or "{}")
    except Exception:
        p = {}
    base = default_params(cfg.preset)
    base.update({k: v for k, v in p.items() if v is not None})
    return base


def set_config(session: Session, section: str, engine: str, preset: str, params: dict | None) -> KbIndexConfig:
    cfg = get_config(session, section)
    cfg.engine = engine if engine in ENGINES else cfg.engine
    cfg.preset = preset
    merged = default_params(preset)
    if params:
        merged.update({k: v for k, v in params.items() if v is not None})
    cfg.params_json = json.dumps(merged)
    cfg.updated_at = _now()
    session.add(cfg)
    session.commit()
    session.refresh(cfg)
    return cfg


# ── Ingestion ────────────────────────────────────────────────────────────────
def add_document(session: Session, section: str, filename: str, content_type: str, data: bytes,
                 seeded: bool = False) -> KbDocument:
    """Parse → chunk → (embed if configured) → store as a KbDocument + KbChunks."""
    text = extract_text(filename, content_type, data)
    cfg = get_config(session, section)
    params = config_params(cfg)
    parts = chunk_text(text, params.get("chunk_size", 1000), params.get("chunk_overlap", 150))

    doc = KbDocument(
        section=section, filename=filename or "document", content_type=content_type or "text/plain",
        size=len(data or b""), text=text, chunk_count=len(parts), indexed=False, seeded=seeded,
    )
    session.add(doc)
    session.commit()
    session.refresh(doc)

    vectors = embeddings.embed(parts) if parts else None
    for i, part in enumerate(parts):
        emb = json.dumps(vectors[i]) if vectors and i < len(vectors) else None
        session.add(KbChunk(doc_id=doc.id, section=section, ordinal=i, text=part, embedding_json=emb))
    doc.indexed = bool(vectors)
    session.add(doc)
    session.commit()
    session.refresh(doc)
    return doc


def delete_document(session: Session, doc_id: int) -> bool:
    doc = session.get(KbDocument, doc_id)
    if not doc:
        return False
    for ch in session.exec(select(KbChunk).where(KbChunk.doc_id == doc_id)).all():
        session.delete(ch)
    session.delete(doc)
    session.commit()
    return True


def reindex(session: Session, section: str) -> int:
    """Recompute embeddings for all chunks in a section. Returns chunks embedded."""
    chunks = session.exec(select(KbChunk).where(KbChunk.section == section)).all()
    if not chunks:
        return 0
    vectors = embeddings.embed([c.text for c in chunks])
    if not vectors:
        # Embeddings unavailable: clear stored vectors, mark docs not-indexed (BM25 still works).
        for c in chunks:
            c.embedding_json = None
            session.add(c)
        for d in session.exec(select(KbDocument).where(KbDocument.section == section)).all():
            d.indexed = False
            session.add(d)
        session.commit()
        return 0
    for c, v in zip(chunks, vectors):
        c.embedding_json = json.dumps(v)
        session.add(c)
    for d in session.exec(select(KbDocument).where(KbDocument.section == section)).all():
        d.indexed = True
        session.add(d)
    session.commit()
    return len(vectors)


# ── Retrieval ────────────────────────────────────────────────────────────────
def _tokenize(s: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", (s or "").lower())


def _bm25_ranking(query: str, texts: list[str]) -> list[float]:
    """BM25 scores per text. Uses rank_bm25 if present, else a light TF fallback."""
    try:
        from rank_bm25 import BM25Okapi  # lazy
        corpus = [_tokenize(t) for t in texts]
        bm = BM25Okapi(corpus)
        return list(bm.get_scores(_tokenize(query)))
    except Exception:
        # Fallback: term-overlap frequency scoring.
        q = set(_tokenize(query))
        scores = []
        for t in texts:
            toks = _tokenize(t)
            if not toks:
                scores.append(0.0)
                continue
            hits = sum(1 for w in toks if w in q)
            scores.append(hits / (len(toks) ** 0.5))
        return scores


def _semantic_ranking(query: str, vectors: list[list[float] | None]) -> list[float] | None:
    """Cosine similarity of query vs each chunk vector; None if not possible."""
    qv = embeddings.embed_one(query)
    if not qv:
        return None
    idx = [i for i, v in enumerate(vectors) if v]
    if not idx:
        return None
    try:
        import numpy as np  # lazy
        mat = np.array([vectors[i] for i in idx], dtype="float32")
        q = np.array(qv, dtype="float32")
        mat /= (np.linalg.norm(mat, axis=1, keepdims=True) + 1e-8)
        q /= (np.linalg.norm(q) + 1e-8)
        sims = mat @ q
        scores = [0.0] * len(vectors)
        for j, i in enumerate(idx):
            scores[i] = float(sims[j])
        return scores
    except Exception as e:
        logger.warning("semantic ranking failed (numpy missing?): %s", e)
        return None


def _rrf(rankings: list[list[int]], k: int = 60) -> dict[int, float]:
    """Reciprocal-rank fusion of several ranked index lists."""
    fused: dict[int, float] = {}
    for ranking in rankings:
        for rank, idx in enumerate(ranking):
            fused[idx] = fused.get(idx, 0.0) + 1.0 / (k + rank + 1)
    return fused


def retrieve(session: Session, section: str, query: str, top_k: int | None = None) -> list[dict]:
    """Return the most relevant chunks for a query in one section.

    Each result: {text, filename, doc_id, score}. Empty list when the section has
    no documents. Engine chosen per admin config, degrading FAISS→BM25 safely.
    """
    chunks = session.exec(select(KbChunk).where(KbChunk.section == section)).all()
    if not chunks:
        return []
    cfg = get_config(session, section)
    params = config_params(cfg)
    k = int(top_k or params.get("top_k", 6))
    min_score = float(params.get("min_score", 0.0))
    engine = cfg.engine

    texts = [c.text for c in chunks]
    vectors = [json.loads(c.embedding_json) if c.embedding_json else None for c in chunks]

    sem = _semantic_ranking(query, vectors) if engine in ("faiss", "hybrid") else None
    bm = _bm25_ranking(query, texts) if engine in ("bm25", "hybrid") or sem is None else None

    if engine == "hybrid" and sem is not None and bm is not None:
        sem_rank = [i for i, _ in sorted(enumerate(sem), key=lambda x: x[1], reverse=True)]
        bm_rank = [i for i, _ in sorted(enumerate(bm), key=lambda x: x[1], reverse=True)]
        fused = _rrf([sem_rank, bm_rank])
        order = sorted(fused, key=lambda i: fused[i], reverse=True)
        scored = [(i, fused[i]) for i in order]
        use_min = 0.0  # RRF scores aren't comparable to cosine thresholds
    elif sem is not None:
        scored = sorted(enumerate(sem), key=lambda x: x[1], reverse=True)
        use_min = min_score
    else:
        scores = bm if bm is not None else _bm25_ranking(query, texts)
        scored = sorted(enumerate(scores), key=lambda x: x[1], reverse=True)
        use_min = 0.0  # BM25 raw scores: keep top_k regardless of an absolute cutoff

    doc_names = {d.id: d.filename for d in session.exec(select(KbDocument).where(KbDocument.section == section)).all()}
    out: list[dict] = []
    for i, score in scored[: max(1, k)]:
        if use_min and score < use_min:
            continue
        c = chunks[i]
        out.append({"text": c.text, "filename": doc_names.get(c.doc_id, "document"),
                    "doc_id": c.doc_id, "score": round(float(score), 4)})
    return out


# ── Status ───────────────────────────────────────────────────────────────────
def status(session: Session) -> dict:
    out = {"embeddings_available": embeddings.is_configured(), "sections": {}}
    for section in KB_SECTIONS:
        docs = session.exec(select(KbDocument).where(KbDocument.section == section)).all()
        cfg = get_config(session, section)
        out["sections"][section] = {
            "documents": len(docs),
            "chunks": sum(d.chunk_count for d in docs),
            "indexed": all(d.indexed for d in docs) if docs else False,
            "engine": cfg.engine,
            "preset": cfg.preset,
            "params": config_params(cfg),
        }
    return out


# ── First-run seed: CREA3 workflow doc ───────────────────────────────────────
CREA3_WORKFLOW_MD = """\
# CREA3 Platform Workflow

CREA3 helps families resolve civil disputes (division of assets in divorce or
inheritance) through a structured, transparent negotiation workflow.

## Stages (in order)

```mermaid
flowchart TD
    A[Create dispute<br/>title + resolution method] --> B[Invite agents<br/>each with entitlement share]
    B --> C{Shares sum to 1.0?}
    C -- no --> B
    C -- yes --> D[List goods<br/>name, value, indivisible?]
    D --> E[Preferences<br/>each party rates goods 1..5 stars]
    E --> F[Strategy / private notes<br/>visible only to mediator]
    F --> G[Ready & validation]
    G --> H[Proposal<br/>suggested allocation]
    H --> I{All parties accept?}
    I -- yes --> K[Final PDF report]
    I -- no --> J[Mediation<br/>schedule a video session]
    J --> H
```

## Field help
- **Entitlement share**: a decimal from 0 to 1 (e.g. 0.5 = half). All parties'
  shares must total 1.0 before the dispute can be validated.
- **Estimated value**: the monetary value of a good, used to measure fairness.
- **Indivisible**: tick when a good cannot be split (e.g. a car).
- **Stars (preferences)**: 1 = don't care, 5 = want it most.

## Roles
- The dispute **owner** manages agents and goods and generates proposals.
- A **mediator** can only view; they cannot submit preferences or proposals.

## What to do next (common)
1. If shares don't sum to 1.0, fix the entitlement shares of the parties.
2. If goods are missing, add them (owner or a joined party).
3. Rate every good you haven't rated yet.
4. Mark yourself ready; once everyone is ready the dispute moves to validation.
5. Review the proposal and accept or decline it.
"""


CREA3_ABOUT_MD = """\
# About the CREA3 Project

**CREA3 — Conflict Resolution with Equitative Algorithms** is an EU co-funded platform
for resolving cross-border civil and consumer disputes online. It brings AI-driven tools
and game-theoretical algorithms to civil dispute resolution, guiding people step by step
toward a fair, efficient and tailored outcome. CREA3 builds on the earlier **CREA** and
**CREA2** projects and is released as open source. It is developed by a European academic
and legal consortium under the **EU Justice (JUST) Programme**.

## European Common Ground of Available Rights (ECGAR)
CREA3 sets aside the mandatory rules of each Member State and operates on the remaining
"available rights", to widen access to Online Dispute Resolution (ODR) and reduce
structural barriers to justice across different national legal systems.

## Three innovations
- **Law & AI standards** — links the ECGAR to the game-theoretical model.
- **Smart conversational interface** — an assistant that guides users through the procedure.
- **Blockchain certification** — smart-contract technology certifies the parties' agreement.

## Consortium & partners
The consortium spans **9 partner institutions across 8 cities in 7 countries** — seven
universities, a continental bar federation, and a consumers' association:

- **Università degli Studi di Napoli Federico II** — Naples, Italy — University — **Project coordinator (lead)** — https://www.unina.it
- **Università degli Studi Suor Orsola Benincasa** — Naples, Italy — University — https://www.unisob.na.it
- **Adiconsum – Associazione Difesa Consumatori e Ambiente** — Rome, Italy — Consumer association — https://www.adiconsum.it
- **Vrije Universiteit Brussel (VUB)** — Brussels, Belgium — University — https://www.vub.be
- **FBE – Fédération des Barreaux d'Europe** — Strasbourg, France — Bar federation — https://www.fbe.org
- **University of Ljubljana** — Ljubljana, Slovenia — University — https://www.uni-lj.si
- **University of Zagreb – Faculty of Law** — Zagreb, Croatia — University — https://www.pravo.unizg.hr
- **Vilnius University** — Vilnius, Lithuania — University — https://www.vu.lt
- **TalTech – Tallinn University of Technology** — Tallinn, Estonia — University — https://www.taltech.ee

The platform targets the six jurisdictions of Italy, Slovenia, Estonia, Belgium, Lithuania
and Croatia. Contact the team at support@crea3.cc.
"""


# Seed docs: filename → content. Each is added only if that filename is absent,
# so this is idempotent and also back-fills new docs into existing databases.
_SEED_DOCS = {
    "CREA3-workflow.md": CREA3_WORKFLOW_MD,
    "CREA3-about.md": CREA3_ABOUT_MD,
}


def seed_workflow_doc(session: Session) -> None:
    """Seed the workflow section with the built-in CREA3 docs (idempotent)."""
    try:
        present = {
            d.filename
            for d in session.exec(select(KbDocument).where(KbDocument.section == "workflow")).all()
        }
        for filename, content in _SEED_DOCS.items():
            if filename in present:
                continue
            add_document(
                session, "workflow", filename, "text/markdown",
                content.encode("utf-8"), seeded=True,
            )
            logger.info("Seeded Knowledge Base 'workflow' section with %s", filename)
    except Exception as e:  # pragma: no cover - never block startup
        logger.warning("KB seed skipped: %s", e)
